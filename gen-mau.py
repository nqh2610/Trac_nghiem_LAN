# Tạo file mau-cau-hoi.docx với hướng dẫn đầy đủ
from docx import Document
from docx.shared import Pt, RGBColor

doc = Document()
doc.styles['Normal'].font.name = 'Times New Roman'
doc.styles['Normal'].font.size = Pt(12)

def th(text, size=13, color=None):
    p = doc.add_paragraph()
    run = p.add_run(text)
    run.bold = True
    run.font.size = Pt(size)
    if color:
        run.font.color.rgb = RGBColor(*color)

def q(text):
    doc.add_paragraph(text)

def blank():
    doc.add_paragraph('')

def code_block(lines):
    doc.add_paragraph('[CODE]')
    for line in lines:
        doc.add_paragraph(line)
    doc.add_paragraph('[/CODE]')

th('MẪU CÂU HỎI - HƯỚNG DẪN SOẠN ĐỀ THI', 14)
th('Hệ thống Trắc Nghiệm LAN', 11)
blank()

# ===== PHẦN 1: CÂU HỎI THÔNG THƯỜNG =====
th('PHẦN 1. CÂU HỎI MỘT ĐÁP ÁN (SINGLE)', 12, (0, 70, 180))
blank()

q('Câu 1. Thủ đô của Việt Nam là thành phố nào?')
q('A. TP. Hồ Chí Minh')
q('B. Hà Nội *')
q('C. Đà Nẵng')
q('D. Cần Thơ')
blank()

q('Câu 2. Năm nào Việt Nam thống nhất đất nước?')
q('A. 1954')
q('B. 1968')
q('C. 1975 *')
q('D. 1986')
blank()

# ===== PHẦN 2: NHIỀU ĐÁP ÁN =====
th('PHẦN 2. CÂU HỎI NHIỀU ĐÁP ÁN - thêm [MULTI] vào đầu câu hỏi', 12, (0, 70, 180))
blank()

q('Câu 3. [MULTI] Chọn các thành phố trực thuộc Trung ương:')
q('A. Hà Nội *')
q('B. Nghệ An')
q('C. TP. Hồ Chí Minh *')
q('D. Đà Nẵng *')
q('E. Cần Thơ *')
blank()

# ===== PHẦN 3: ĐÚNG/SAI =====
th('PHẦN 3. CÂU HỎI ĐÚNG/SAI - thêm [DUNG/SAI] vào đầu câu hỏi', 12, (0, 70, 180))
blank()

q('Câu 4. [DUNG/SAI] Xác định tính đúng/sai:')
q('A. Hà Nội là thủ đô Việt Nam. (đúng)')
q('B. Việt Nam không có biên giới với Trung Quốc. (sai)')
q('C. Sông Hồng chảy qua Hà Nội. (đúng)')
q('D. Việt Nam có 63 tỉnh thành. (đúng)')
blank()

# ===== PHẦN 4: CÔNG THỨC TOÁN =====
th('PHẦN 4. CÔNG THỨC TOÁN - dùng \\( ... \\) cho LaTeX', 12, (180, 0, 0))
blank()

p = doc.add_paragraph()
p.add_run('Cú pháp: ').bold = True
p.add_run('Bọc công thức LaTeX trong \\( ... \\) cho inline, hoặc \\[ ... \\] cho dòng riêng')
blank()

q('Câu 5. Giá trị của \\(\\sqrt{4} + \\sqrt{9}\\) bằng bao nhiêu?')
q('A. 2')
q('B. 3')
q('C. 5 *')
q('D. 7')
blank()

q('Câu 6. Nghiệm của phương trình \\(x^2 - 5x + 6 = 0\\) là:')
q('A. \\(x = 1\\) hoặc \\(x = 6\\)')
q('B. \\(x = 2\\) hoặc \\(x = 3\\) *')
q('C. \\(x = -2\\) hoặc \\(x = -3\\)')
q('D. \\(x = 0\\) hoặc \\(x = 5\\)')
blank()

q('Câu 7. [DUNG/SAI] Cho hàm số \\(f(x) = x^2 - 4\\), xác định đúng/sai:')
q('A. \\(f(0) = -4\\) (đúng)')
q('B. \\(f(2) = 0\\) (đúng)')
q('C. \\(f(-2) = 4\\) (sai)')
q('D. Hàm số đạt cực tiểu tại \\(x = 0\\) (đúng)')
blank()

p2 = doc.add_paragraph()
p2.add_run('Các ký hiệu LaTeX thường dùng:').bold = True
doc.add_paragraph('\\(\\frac{a}{b}\\)  -->  phân số a/b')
doc.add_paragraph('\\(\\sqrt{x}\\)  -->  căn bậc hai')
doc.add_paragraph('\\(\\sqrt[3]{x}\\)  -->  căn bậc ba')
doc.add_paragraph('\\(x^{2}\\)  -->  lũy thừa')
doc.add_paragraph('\\(\\sum_{i=1}^{n} i\\)  -->  tổng sigma')
doc.add_paragraph('\\(\\int_{a}^{b} f(x)dx\\)  -->  tích phân')
doc.add_paragraph('\\(\\lim_{x \\to 0}\\)  -->  giới hạn')
doc.add_paragraph('\\(\\pi, \\alpha, \\beta, \\theta\\)  -->  chữ Hy Lạp')
blank()

# ===== PHẦN 5: CODE =====
th('PHẦN 5. HIỂN THỊ CODE - dùng [CODE] và [/CODE]', 12, (180, 0, 0))
blank()

p3 = doc.add_paragraph()
p3.add_run('Cú pháp: ').bold = True
p3.add_run('Đặt [CODE] trên một dòng riêng, viết code, rồi đặt [/CODE] trên dòng riêng')
blank()

q('Câu 8. Xét đoạn code Python sau, kết quả in ra là gì?')
code_block([
    'x = 10',
    'y = 3',
    'print(x // y)',
])
q('A. 3 *')
q('B. 3.33')
q('C. 4')
q('D. 1')
blank()

q('Câu 9. Đoạn code C++ sau in ra giá trị gì?')
code_block([
    '#include <iostream>',
    'using namespace std;',
    'int main() {',
    '    int x = 5;',
    '    cout << x * x;',
    '}',
])
q('A. 5')
q('B. 10')
q('C. 25 *')
q('D. 55')
blank()

q('Câu 10. [MULTI] Xét đoạn HTML sau, khẳng định nào đúng?')
code_block([
    '<div class="box">',
    '  <p>Hello World</p>',
    '</div>',
])
q('A. Thẻ div chứa thẻ p *')
q('B. Thuộc tính class có giá trị "box" *')
q('C. Thẻ p không cần đóng')
q('D. Đây là HTML5 hợp lệ *')
blank()

doc.save('data/mau-cau-hoi.docx')
print('OK: data/mau-cau-hoi.docx')
